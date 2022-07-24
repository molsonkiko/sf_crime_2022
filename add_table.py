from openpyxl.worksheet.table import Table, TableStyleInfo
from excel_ranges import range_containing_table
import openpyxl


def add_table(fname, sheet_name, tbl_name, rng, force=False):
    '''For an Excel workbook at path fname, add a table named tbl_name to 
    range rng of sheet sheet_name.
rng should be an Excel range of the form 'A1:B20', or a pandas DataFrame.
    If rng is a pandas DataFrame, the table will automatically have its upper
    left corner at cell A1 of the sheet.
If you want to customize the location of a table, use the 
    range_containing_table function from excel_ranges.
If force is True and a table already exists with tbl_name, remove that table and replace it with the new one
    '''
    wb = openpyxl.open(fname)
    if 'DataFrame' in str(type(rng)):
        rng = range_containing_table(rng)
    try:
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        tbl = Table(displayName=tbl_name, ref=rng)
        tbl.tableStyleInfo = style
        if force and tbl_name in wb[sheet_name].tables:
            del wb[sheet_name].tables[tbl_name]
        wb[sheet_name].add_table(tbl)
    finally:
        wb.save(fname)
        wb.close()
